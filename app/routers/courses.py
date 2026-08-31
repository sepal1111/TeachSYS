import io
import csv
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import urllib.parse
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Response
from sqlite3 import Connection
from typing import List, Optional

from app.database import get_db
from app.models import CourseCreate, CourseUpdate, StudentCreate, StudentBatchImportRequest, TextImportRequest

router = APIRouter(prefix="/api/courses", tags=["Courses"])

DEFAULT_RULES = [
    {"title": "熱心助人", "category": "positive", "score_value": 1, "icon": "🤝"},
    {"title": "發言踴躍", "category": "positive", "score_value": 1, "icon": "🙋‍♂️"},
    {"title": "專心聽講", "category": "positive", "score_value": 1, "icon": "👂"},
    {"title": "作業優良", "category": "positive", "score_value": 2, "icon": "📝"},
    {"title": "團隊合作", "category": "positive", "score_value": 1, "icon": "🌟"},
    {"title": "上課吵鬧", "category": "negative", "score_value": -1, "icon": "📢"},
    {"title": "未帶用品", "category": "negative", "score_value": -1, "icon": "🎒"},
    {"title": "上課分心", "category": "negative", "score_value": -1, "icon": "😴"},
]

@router.get("")
def list_courses(db: Connection = Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("""
        SELECT c.*, 
               (SELECT COUNT(*) FROM students s WHERE s.course_id = c.id AND s.is_active = 1) AS student_count
        FROM courses c
        ORDER BY c.id DESC
    """)
    courses = [dict(row) for row in cursor.fetchall()]
    return courses

@router.post("")
def create_course(data: CourseCreate, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    cursor.execute(
        "INSERT INTO courses (name, teacher_type, description) VALUES (?, ?, ?)",
        (data.name, data.teacher_type, data.description)
    )
    course_id = cursor.lastrowid
    
    # Seed default evaluation rules
    for rule in DEFAULT_RULES:
        cursor.execute("""
            INSERT INTO evaluation_rules (course_id, title, category, score_value, icon, is_default)
            VALUES (?, ?, ?, ?, ?, 1)
        """, (course_id, rule["title"], rule["category"], rule["score_value"], rule["icon"]))
        
    db.commit()
    return {"id": course_id, "message": "Course created with default evaluation rules"}

@router.get("/{course_id}")
def get_course(course_id: int, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("SELECT * FROM courses WHERE id = ?", (course_id,))
    row = cursor.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Course not found")
    return dict(row)

@router.delete("/{course_id}")
def delete_course(course_id: int, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("DELETE FROM courses WHERE id = ?", (course_id,))
    db.commit()
    return {"message": "Course deleted successfully"}

from app.models import CourseCreate, CourseUpdate, StudentCreate, StudentUpdate, StudentBatchImportRequest, TextImportRequest, EvaluationRuleCreate

@router.get("/{course_id}/rules")
def list_course_rules(course_id: int, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("""
        SELECT * FROM evaluation_rules WHERE course_id = ? ORDER BY category DESC, id ASC
    """, (course_id,))
    return [dict(row) for row in cursor.fetchall()]

@router.post("/{course_id}/rules")
def create_course_rule(course_id: int, data: EvaluationRuleCreate, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("""
        INSERT INTO evaluation_rules (course_id, title, category, score_value, icon)
        VALUES (?, ?, ?, ?, ?)
    """, (course_id, data.title, data.category, data.score_value, data.icon))
    db.commit()
    return {"id": cursor.lastrowid, "message": "Rule created"}

@router.put("/{course_id}/rules/{rule_id}")
def update_course_rule(course_id: int, rule_id: int, data: EvaluationRuleCreate, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("""
        UPDATE evaluation_rules
        SET title = ?, category = ?, score_value = ?, icon = ?
        WHERE id = ? AND course_id = ?
    """, (data.title, data.category, data.score_value, data.icon, rule_id, course_id))
    db.commit()
    return {"message": "Rule updated"}

@router.delete("/{course_id}/rules/{rule_id}")
def delete_course_rule(course_id: int, rule_id: int, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("DELETE FROM evaluation_rules WHERE id = ? AND course_id = ?", (rule_id, course_id))
    db.commit()
    return {"message": "Rule deleted"}

@router.post("/{course_id}/rules/reset_defaults")
def reset_course_default_rules(course_id: int, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("DELETE FROM evaluation_rules WHERE course_id = ?", (course_id,))
    for rule in DEFAULT_RULES:
        cursor.execute("""
            INSERT INTO evaluation_rules (course_id, title, category, score_value, icon)
            VALUES (?, ?, ?, ?, ?)
        """, (course_id, rule["title"], rule["category"], rule["score_value"], rule["icon"]))
    db.commit()
    return {"message": "Default rules restored"}

# --- 學生管理 ---

@router.get("/{course_id}/students")
def list_students(course_id: int, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("""
        SELECT s.*, g.group_name 
        FROM students s
        LEFT JOIN groups g ON s.group_id = g.id
        WHERE s.course_id = ? AND s.is_active = 1
        ORDER BY s.student_number ASC
    """, (course_id,))
    students = [dict(row) for row in cursor.fetchall()]
    return students

@router.post("/{course_id}/students")
def create_student(course_id: int, data: StudentCreate, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    cursor.execute(
        "INSERT INTO students (course_id, student_number, name, english_name, gender, group_id, student_code) VALUES (?, ?, ?, ?, ?, ?, ?)",
        (course_id, data.student_number, data.name, data.english_name, data.gender, data.group_id, data.student_code)
    )
    db.commit()
    return {"id": cursor.lastrowid, "message": "Student created"}

@router.post("/{course_id}/students/batch")
def batch_import_students(course_id: int, data: StudentBatchImportRequest, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    count = 0
    for item in data.students:
        gender = (item.gender or "M").upper()
        if gender not in ["M", "F"]:
            gender = "M"
        cursor.execute(
            "INSERT INTO students (course_id, student_number, name, english_name, gender, student_code) VALUES (?, ?, ?, ?, ?, ?)",
            (course_id, item.student_number, item.name, item.english_name, gender, item.student_code)
        )
        count += 1
    db.commit()
    return {"imported_count": count, "message": f"Successfully imported {count} students"}

@router.post("/{course_id}/students/text_import")
def text_import_students(course_id: int, data: TextImportRequest, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    lines = data.text_content.strip().split("\n")
    imported_students = []
    
    auto_num = 1
    for line in lines:
        line = line.strip()
        if not line:
            continue
        parts = [p.strip() for p in line.replace(",", " ").replace("\t", " ").split() if p.strip()]
        if not parts:
            continue
        
        num = None
        code = None
        name = ""
        english_name = None
        gender = "M"

        if parts[0].isdigit():
            num = int(parts[0])
            if len(parts) >= 5:
                # Format: 座號 學號 中文姓名 英文姓名 性別
                code = parts[1]
                name = parts[2]
                english_name = parts[3]
                if "女" in parts[4].upper() or parts[4].upper() == "F":
                    gender = "F"
            elif len(parts) == 4:
                # Format could be: 座號 學號 姓名 性別 OR 座號 姓名 英文名 性別
                last_is_gender = "女" in parts[3].upper() or parts[3].upper() in ["F", "M", "男"]
                if last_is_gender:
                    if "女" in parts[3].upper() or parts[3].upper() == "F":
                        gender = "F"
                    # Check if part[1] is numeric code
                    if parts[1].isdigit() and len(parts[1]) >= 4:
                        code = parts[1]
                        name = parts[2]
                    else:
                        name = parts[1]
                        english_name = parts[2]
                else:
                    code = parts[1]
                    name = parts[2]
                    english_name = parts[3]
            elif len(parts) == 3:
                # Format: 座號 姓名 性別 OR 座號 姓名 英文名
                name = parts[1]
                if "女" in parts[2].upper() or parts[2].upper() == "F":
                    gender = "F"
                elif "男" in parts[2].upper() or parts[2].upper() == "M":
                    gender = "M"
                else:
                    english_name = parts[2]
            else:
                name = parts[1] if len(parts) > 1 else f"學生{num}"
        else:
            num = auto_num
            name = parts[0]
            if len(parts) >= 3:
                english_name = parts[1]
                if "女" in parts[2].upper() or parts[2].upper() == "F":
                    gender = "F"
            elif len(parts) >= 2:
                if "女" in parts[1].upper() or parts[1].upper() == "F":
                    gender = "F"
                else:
                    english_name = parts[1]

        auto_num = max(auto_num, num + 1)
        imported_students.append({
            "student_number": num,
            "student_code": code,
            "name": name,
            "english_name": english_name,
            "gender": gender
        })

    count = 0
    for s in imported_students:
        cursor.execute(
            "INSERT INTO students (course_id, student_number, student_code, name, english_name, gender) VALUES (?, ?, ?, ?, ?, ?)",
            (course_id, s["student_number"], s["student_code"], s["name"], s["english_name"], s["gender"])
        )
        count += 1
    db.commit()
    return {"imported_count": count, "message": f"Successfully imported {count} students"}

@router.post("/{course_id}/students/upload")
async def upload_students_file(course_id: int, file: UploadFile = File(...), db: Connection = Depends(get_db)):
    filename = file.filename.lower()
    contents = await file.read()
    imported_students = []

    if filename.endswith(".csv"):
        text = contents.decode("utf-8-sig", errors="ignore")
        reader = csv.reader(io.StringIO(text))
        for row in reader:
            if not row or len(row) < 2:
                continue
            if "座號" in row[0] or "姓名" in row[1] or not row[0].strip().isdigit():
                continue
            num = int(row[0].strip())
            code = None
            name = ""
            english_name = None
            gender = "M"

            row_str = [cell.strip() for cell in row]
            if len(row_str) >= 5:
                # 座號, 學號, 中文姓名, 英文姓名, 性別
                code = row_str[1]
                name = row_str[2]
                english_name = row_str[3] or None
                g_str = row_str[4].upper()
                if "女" in g_str or g_str == "F":
                    gender = "F"
            elif len(row_str) == 4:
                # Check if col 3 is gender or english name
                g_str = row_str[3].upper()
                if "女" in g_str or g_str == "F":
                    gender = "F"
                    code = row_str[1]
                    name = row_str[2]
                elif "男" in g_str or g_str == "M":
                    gender = "M"
                    code = row_str[1]
                    name = row_str[2]
                else:
                    # 座號, 中文姓名, 英文姓名, 性別
                    name = row_str[1]
                    english_name = row_str[2] or None
                    if "女" in row_str[3].upper() or row_str[3].upper() == "F":
                        gender = "F"
            elif len(row_str) == 3:
                name = row_str[1]
                g_str = row_str[2].upper()
                if "女" in g_str or g_str == "F":
                    gender = "F"
            else:
                name = row_str[1]

            imported_students.append({
                "student_number": num,
                "student_code": code,
                "name": name,
                "english_name": english_name,
                "gender": gender
            })

    elif filename.endswith(".xlsx") or filename.endswith(".xls"):
        wb = openpyxl.load_workbook(filename=io.BytesIO(contents))
        sheet = wb.active
        for row in sheet.iter_rows(values_only=True):
            if not row or len(row) < 2:
                continue
            val0 = str(row[0]).strip() if row[0] is not None else ""
            if not val0.isdigit():
                continue
            num = int(val0)
            code = None
            name = ""
            english_name = None
            gender = "M"

            row_str = [str(cell).strip() if cell is not None else "" for cell in row]

            if len(row_str) >= 5:
                code = row_str[1]
                name = row_str[2]
                english_name = row_str[3] or None
                g_str = row_str[4].upper()
                if "女" in g_str or g_str == "F":
                    gender = "F"
            elif len(row_str) == 4:
                g_str = row_str[3].upper()
                if "女" in g_str or g_str == "F":
                    gender = "F"
                    code = row_str[1]
                    name = row_str[2]
                elif "男" in g_str or g_str == "M":
                    gender = "M"
                    code = row_str[1]
                    name = row_str[2]
                else:
                    name = row_str[1]
                    english_name = row_str[2] or None
                    if "女" in row_str[3].upper() or row_str[3].upper() == "F":
                        gender = "F"
            elif len(row_str) == 3:
                name = row_str[1]
                g_str = row_str[2].upper()
                if "女" in g_str or g_str == "F":
                    gender = "F"
            else:
                name = row_str[1]

            imported_students.append({
                "student_number": num,
                "student_code": code,
                "name": name,
                "english_name": english_name,
                "gender": gender
            })
    else:
        raise HTTPException(status_code=400, detail="Only .csv and .xlsx/.xls files are supported")

    cursor = db.cursor()
    count = 0
    for s in imported_students:
        cursor.execute(
            "INSERT INTO students (course_id, student_number, student_code, name, english_name, gender) VALUES (?, ?, ?, ?, ?, ?)",
            (course_id, s["student_number"], s["student_code"], s["name"], s["english_name"], s["gender"])
        )
        count += 1
    db.commit()
    return {"imported_count": count, "message": f"Successfully imported {count} students"}

@router.put("/{course_id}/students/{student_id}")
def update_student(course_id: int, student_id: int, data: StudentUpdate, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("SELECT * FROM students WHERE id = ? AND course_id = ?", (student_id, course_id))
    student = cursor.fetchone()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")

    num = data.student_number if data.student_number is not None else student["student_number"]
    code = data.student_code if data.student_code is not None else student["student_code"]
    name = data.name if data.name is not None else student["name"]
    english_name = data.english_name if data.english_name is not None else student["english_name"]
    gender = data.gender if data.gender is not None else student["gender"]

    cursor.execute("""
        UPDATE students
        SET student_number = ?, student_code = ?, name = ?, english_name = ?, gender = ?
        WHERE id = ? AND course_id = ?
    """, (num, code, name, english_name, gender, student_id, course_id))
    db.commit()
    return {"message": "Student updated successfully"}

@router.delete("/{course_id}/students/{student_id}")
def delete_student(course_id: int, student_id: int, db: Connection = Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("DELETE FROM students WHERE id = ? AND course_id = ?", (student_id, course_id))
    db.commit()
    return {"message": "Student deleted"}

# --- 範本檔案下載 ---

@router.get("/template/students_excel")
def download_student_excel_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "學生名冊匯入範例"

    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_font = Font(name="Microsoft JhengHei", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Microsoft JhengHei", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = ["座號", "學號", "中文姓名", "英文姓名", "性別"]
    ws.append(headers)

    sample_data = [
        [1, "112001", "王小明", "David", "男"],
        [2, "112002", "李小華", "Emily", "女"],
        [3, "112003", "張大同", "Tom", "男"],
        [4, "112004", "陳雅婷", "Grace", "女"],
        [5, "112005", "林志豪", "Leo", "男"],
        [6, "112006", "黃美玲", "May", "女"],
        [7, "112007", "趙子龍", "Alex", "男"],
        [8, "112008", "周雅玲", "Chloe", "女"],
        [9, "112009", "孫悟空", "Sam", "男"],
        [10, "112010", "吳小雯", "Wendy", "女"]
    ]

    for row in sample_data:
        ws.append(row)

    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 12

    for col_idx in range(1, 6):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for r in range(2, len(sample_data) + 2):
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.font = cell_font
            cell.alignment = center_align
            cell.border = thin_border

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = "學生名冊匯入範例.xlsx"
    encoded_filename = urllib.parse.quote(filename)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"
        }
    )

@router.get("/template/students_csv")
def download_student_csv_template():
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["座號", "學號", "中文姓名", "英文姓名", "性別"])
    sample_data = [
        [1, "112001", "王小明", "David", "男"],
        [2, "112002", "李小華", "Emily", "女"],
        [3, "112003", "張大同", "Tom", "男"],
        [4, "112004", "陳雅婷", "Grace", "女"],
        [5, "112005", "林志豪", "Leo", "男"]
    ]
    for row in sample_data:
        writer.writerow(row)

    # Encode with UTF-8 BOM so Excel opens it with correct traditional Chinese encoding
    content_bytes = "\ufeff".encode("utf-8") + output.getvalue().encode("utf-8")
    filename = "學生名冊匯入範例.csv"
    encoded_filename = urllib.parse.quote(filename)

    return Response(
        content=content_bytes,
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"
        }
    )
