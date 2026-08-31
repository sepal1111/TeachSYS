import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import date, timedelta
import calendar
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlite3 import Connection
from typing import Optional

from app.database import get_db
from app.timezone import get_today_taipei, get_today_str_taipei, get_now_str_taipei

router = APIRouter(prefix="/api/reports", tags=["Reports & Dashboard"])

@router.get("/{course_id}/dashboard")
def get_dashboard_data(
    course_id: int,
    period: str = Query("today", description="today | week | month | range | semester"),
    plan_id: Optional[int] = Query(None, description="指定分組模式 ID (預設為當前生效模式)"),
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    db: Connection = Depends(get_db)
):
    today = get_today_taipei()
    today_str = today.isoformat()

    if period == "today":
        start_date = today_str
        end_date = today_str
    elif period == "week":
        monday = today - timedelta(days=today.weekday())
        sunday = monday + timedelta(days=6)
        start_date = monday.isoformat()
        end_date = sunday.isoformat()
    elif period == "month":
        first_day = date(today.year, today.month, 1)
        last_day_num = calendar.monthrange(today.year, today.month)[1]
        last_day = date(today.year, today.month, last_day_num)
        start_date = first_day.isoformat()
        end_date = last_day.isoformat()
    elif period == "semester":
        start_date = None
        end_date = None

    cursor = db.cursor()

    # 0. Fetch active/specified plan
    if isinstance(plan_id, int):
        cursor.execute("SELECT id, name, is_active FROM group_plans WHERE id = ? AND course_id = ?", (plan_id, course_id))
        active_plan = cursor.fetchone()
    else:
        cursor.execute("SELECT id, name, is_active FROM group_plans WHERE course_id = ? AND is_active = 1 LIMIT 1", (course_id,))
        active_plan = cursor.fetchone()

    if not active_plan:
        cursor.execute("SELECT id, name, is_active FROM group_plans WHERE course_id = ? ORDER BY id ASC LIMIT 1", (course_id,))
        active_plan = cursor.fetchone()

    plan_id_val = active_plan["id"] if active_plan else None

    # 1. Fetch all active students with group_name from group_members for active_plan
    if plan_id_val:
        cursor.execute("""
            SELECT s.id, s.student_number, s.student_code, s.name, s.english_name, s.gender,
                   gm.group_id, g.group_name, g.icon_url AS group_icon_url
            FROM students s
            LEFT JOIN group_members gm ON s.id = gm.student_id AND gm.plan_id = ?
            LEFT JOIN groups g ON gm.group_id = g.id
            WHERE s.course_id = ? AND s.is_active = 1
            ORDER BY s.student_number ASC
        """, (plan_id_val, course_id))
    else:
        cursor.execute("""
            SELECT s.id, s.student_number, s.student_code, s.name, s.english_name, s.gender,
                   s.group_id, g.group_name, g.icon_url AS group_icon_url
            FROM students s
            LEFT JOIN groups g ON s.group_id = g.id
            WHERE s.course_id = ? AND s.is_active = 1
            ORDER BY s.student_number ASC
        """, (course_id,))
    students = [dict(r) for r in cursor.fetchall()]

    # 2. Fetch today's attendance for visual disabling indicator
    cursor.execute("""
        SELECT student_id, status FROM attendance WHERE course_id = ? AND date = ?
    """, (course_id, today_str))
    attendance_map = {row["student_id"]: row["status"] for row in cursor.fetchall()}

    # 3. Calculate score per student in period (all score logs for personal cumulative total)
    score_query = """
        SELECT student_id, SUM(score) AS total_score
        FROM score_logs
        WHERE course_id = ? AND is_undone = 0
    """
    params = [course_id]

    if period == "range" and not start_date and not end_date:
        score_query += " AND 1 = 0"
    else:
        if start_date:
            score_query += " AND date >= ?"
            params.append(start_date)
        if end_date:
            score_query += " AND date <= ?"
            params.append(end_date)

    score_query += " GROUP BY student_id"
    cursor.execute(score_query, params)
    score_map = {row["student_id"]: row["total_score"] for row in cursor.fetchall()}

    # 4. Calculate isolated group award scores for each group in this plan
    # (Each group award event counts as the award score directly, e.g. +2 points, not multiplied by student count)
    grp_score_query = """
        SELECT group_id, SUM(event_score) AS group_score
        FROM (
            SELECT sl.group_id, sl.timestamp, sl.rule_title, sl.score AS event_score
            FROM score_logs sl
            LEFT JOIN groups g ON sl.group_id = g.id
            WHERE sl.course_id = ? AND sl.is_undone = 0 AND sl.group_id IS NOT NULL
    """
    grp_params = [course_id]
    if plan_id_val:
        grp_score_query += " AND (sl.plan_id = ? OR (sl.plan_id IS NULL AND (g.plan_id = ? OR g.plan_id IS NULL)))"
        grp_params.extend([plan_id_val, plan_id_val])

    if period == "range" and not start_date and not end_date:
        grp_score_query += " AND 1 = 0"
    else:
        if start_date:
            grp_score_query += " AND sl.date >= ?"
            grp_params.append(start_date)
        if end_date:
            grp_score_query += " AND sl.date <= ?"
            grp_params.append(end_date)

    grp_score_query += """
            GROUP BY sl.group_id, sl.timestamp, sl.rule_title, sl.score
        )
        GROUP BY group_id
    """
    cursor.execute(grp_score_query, grp_params)
    group_award_map = {row["group_id"]: row["group_score"] for row in cursor.fetchall()}

    # 5. Fetch groups defined in this plan
    if plan_id_val:
        cursor.execute("SELECT id, group_name, icon_url FROM groups WHERE course_id = ? AND plan_id = ? ORDER BY order_index ASC, id ASC", (course_id, plan_id_val))
        plan_groups = [dict(r) for r in cursor.fetchall()]
    else:
        cursor.execute("SELECT id, group_name, icon_url FROM groups WHERE course_id = ? ORDER BY order_index ASC, id ASC", (course_id,))
        plan_groups = [dict(r) for r in cursor.fetchall()]

    member_count_map = {}
    for s in students:
        s["score"] = score_map.get(s["id"], 0)
        s["today_attendance"] = attendance_map.get(s["id"], "present")
        s["is_absent"] = s["today_attendance"] in ["absent", "sick_leave", "personal_leave", "official_leave", "bereavement_leave"]
        gid = s.get("group_id")
        if gid:
            member_count_map[gid] = member_count_map.get(gid, 0) + 1

    group_scores = []
    for g in plan_groups:
        gid = g["id"]
        total_pts = group_award_map.get(gid, 0)
        m_count = member_count_map.get(gid, 0)
        group_scores.append({
            "id": gid,
            "group_name": g["group_name"],
            "icon_url": g.get("icon_url"),
            "total_score": total_pts,
            "member_count": m_count,
            "avg_score": total_pts
        })

    # Individual Leaderboard (Only rank students with score > 0)
    scored_students = [s for s in students if s["score"] > 0]
    individual_leaderboard = sorted(scored_students, key=lambda x: x["score"], reverse=True)

    # Group Leaderboard (Only rank groups with total_score > 0)
    scored_groups = [g for g in group_scores if g["total_score"] > 0]
    group_leaderboard = sorted(scored_groups, key=lambda x: x["total_score"], reverse=True)

    return {
        "period": period,
        "plan": dict(active_plan) if active_plan else None,
        "start_date": start_date,
        "end_date": end_date,
        "students": students,
        "individual_leaderboard": individual_leaderboard[:10],
        "group_leaderboard": group_leaderboard
    }

@router.get("/{course_id}/export")
def export_excel_report(
    course_id: int,
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    db: Connection = Depends(get_db)
):
    cursor = db.cursor()
    cursor.execute("SELECT name FROM courses WHERE id = ?", (course_id,))
    course_row = cursor.fetchone()
    if not course_row:
        raise HTTPException(status_code=404, detail="Course not found")
    course_name = course_row["name"]

    wb = openpyxl.Workbook()
    # Style definitions
    header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Arial", size=14, bold=True)
    border_side = Side(style='thin', color='D1D5DB')
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    # --- Sheet 1: 出缺席統計 ---
    ws_att = wb.active
    ws_att.title = "出缺席統計"
    ws_att.append([f"課程：{course_name} - 出缺席統計報表", f"區間：{start_date or '不限'} ~ {end_date or '不限'}"])
    ws_att.append([])

    headers_att = ["座號", "姓名", "性別", "出席次數", "病假次數", "事假次數", "公假次數", "喪假次數", "遲到次數"]
    ws_att.append(headers_att)

    cursor.execute("""
        SELECT id, student_number, name, gender FROM students WHERE course_id = ? AND is_active = 1 ORDER BY student_number ASC
    """, (course_id,))
    students = [dict(r) for r in cursor.fetchall()]

    for s in students:
        att_query = "SELECT status, COUNT(*) AS cnt FROM attendance WHERE student_id = ?"
        att_params = [s["id"]]
        if start_date:
            att_query += " AND date >= ?"
            att_params.append(start_date)
        if end_date:
            att_query += " AND date <= ?"
            att_params.append(end_date)
        att_query += " GROUP BY status"
        
        cursor.execute(att_query, att_params)
        counts = {r["status"]: r["cnt"] for r in cursor.fetchall()}

        ws_att.append([
            s["student_number"],
            s["name"],
            "男" if s["gender"] == "M" else "女",
            counts.get("present", 0),
            counts.get("sick_leave", 0) + counts.get("absent", 0),
            counts.get("personal_leave", 0),
            counts.get("official_leave", 0),
            counts.get("bereavement_leave", 0),
            counts.get("late", 0)
        ])

    # --- Sheet 2: 量化成績計分表 ---
    ws_score = wb.create_sheet(title="量化成績統計與明細")
    ws_score.append([f"課程：{course_name} - 量化成績總計與評分明細"])
    ws_score.append([])
    ws_score.append(["座號", "姓名", "加減分總計"])

    score_log_query = """
        SELECT sl.date, sl.timestamp, s.student_number, s.name AS student_name, sl.rule_title, sl.score, sl.category,
               gp.name AS plan_name, g.group_name
        FROM score_logs sl
        JOIN students s ON sl.student_id = s.id
        LEFT JOIN group_plans gp ON sl.plan_id = gp.id
        LEFT JOIN groups g ON sl.group_id = g.id
        WHERE sl.course_id = ? AND sl.is_undone = 0
    """
    score_log_params = [course_id]

    if start_date:
        score_log_query += " AND sl.date >= ?"
        score_log_params.append(start_date)
    if end_date:
        score_log_query += " AND sl.date <= ?"
        score_log_params.append(end_date)

    for s in students:
        # Sum score
        sum_query = "SELECT SUM(score) AS total FROM score_logs WHERE student_id = ? AND is_undone = 0"
        sum_params = [s["id"]]
        if start_date:
            sum_query += " AND date >= ?"
            sum_params.append(start_date)
        if end_date:
            sum_query += " AND date <= ?"
            sum_params.append(end_date)
        cursor.execute(sum_query, sum_params)
        total_val = cursor.fetchone()["total"] or 0

        ws_score.append([s["student_number"], s["name"], total_val])

    ws_score.append([])
    ws_score.append(["--- 評分細項明細 ---"])
    ws_score.append(["日期時間", "座號", "學生姓名", "評分項目", "分數", "類別", "所屬分組模式", "所屬小組"])

    score_log_query += " ORDER BY sl.timestamp DESC"
    cursor.execute(score_log_query, score_log_params)
    for log in cursor.fetchall():
        ws_score.append([
            log["timestamp"],
            log["student_number"],
            log["student_name"],
            log["rule_title"],
            log["score"],
            "正向" if log["category"] == "positive" else "負向",
            log["plan_name"] or "個人評分",
            log["group_name"] or "-"
        ])

    # --- Sheet 3: 小組分數成員與得分一覽 ---
    ws_groups = wb.create_sheet(title="小組分數成員與得分一覽")
    ws_groups.append([f"課程：{course_name} - 小組分組名單、獨立總得分與成員一覽", f"區間：{start_date or '不限'} ~ {end_date or '不限'}"])
    ws_groups.append([])

    cursor.execute("SELECT id, name, is_active FROM group_plans WHERE course_id = ? ORDER BY id ASC", (course_id,))
    plans = [dict(r) for r in cursor.fetchall()]

    all_plan_group_data = []
    max_members = 1

    pastel_colors = [
        "EFF6FF",  # 淡藍 (Soft Blue)
        "ECFDF5",  # 淡薄荷綠 (Soft Mint Green)
        "FFFBEB",  # 淡杏黃 (Soft Amber)
        "FAF5FF",  # 淡紫丁香 (Soft Lavender)
        "FFF1F2",  # 淡玫瑰粉 (Soft Rose)
        "F0FDFA",  # 淡青綠 (Soft Teal)
        "FFF7ED",  # 淡蜜桃 (Soft Peach)
        "F5F3FF",  # 淡紫藤 (Soft Violet)
    ]
    pastel_fills = [PatternFill(start_color=c, end_color=c, fill_type="solid") for c in pastel_colors]

    for plan_idx, p in enumerate(plans):
        pid = p["id"]
        pname = p["name"]

        cursor.execute("SELECT id, group_name FROM groups WHERE course_id = ? AND plan_id = ? ORDER BY order_index ASC, id ASC", (course_id, pid))
        p_groups = [dict(r) for r in cursor.fetchall()]

        p_grp_score_query = """
            SELECT group_id, SUM(event_score) AS group_score
            FROM (
                SELECT sl.group_id, sl.timestamp, sl.rule_title, sl.score AS event_score
                FROM score_logs sl
                LEFT JOIN groups g ON sl.group_id = g.id
                WHERE sl.course_id = ? AND sl.is_undone = 0 AND sl.group_id IS NOT NULL
                  AND (sl.plan_id = ? OR (sl.plan_id IS NULL AND (g.plan_id = ? OR g.plan_id IS NULL)))
        """
        p_grp_params = [course_id, pid, pid]
        if start_date:
            p_grp_score_query += " AND sl.date >= ?"
            p_grp_params.append(start_date)
        if end_date:
            p_grp_score_query += " AND sl.date <= ?"
            p_grp_params.append(end_date)
        p_grp_score_query += " GROUP BY sl.group_id, sl.timestamp, sl.rule_title, sl.score) GROUP BY group_id"
        cursor.execute(p_grp_score_query, p_grp_params)
        p_group_score_map = {row["group_id"]: row["group_score"] for row in cursor.fetchall()}

        for g in p_groups:
            gid = g["id"]
            gscore = p_group_score_map.get(gid, 0)
            cursor.execute("""
                SELECT s.student_number, s.name
                FROM group_members gm
                JOIN students s ON gm.student_id = s.id
                WHERE gm.plan_id = ? AND gm.group_id = ? AND s.is_active = 1
                ORDER BY s.student_number ASC
            """, (pid, gid))
            members = cursor.fetchall()
            if len(members) > max_members:
                max_members = len(members)
            all_plan_group_data.append({
                "plan_idx": plan_idx,
                "plan_name": pname,
                "group_name": g["group_name"],
                "group_score": gscore,
                "member_count": len(members),
                "members": [f"{m['student_number']}號 {m['name']}" for m in members]
            })

        # Unassigned members in this plan
        cursor.execute("""
            SELECT s.student_number, s.name
            FROM students s
            WHERE s.course_id = ? AND s.is_active = 1
              AND s.id NOT IN (
                  SELECT student_id FROM group_members WHERE plan_id = ? AND group_id IS NOT NULL AND group_id > 0
              )
            ORDER BY s.student_number ASC
        """, (course_id, pid))
        unassigned_members = cursor.fetchall()
        if unassigned_members:
            if len(unassigned_members) > max_members:
                max_members = len(unassigned_members)
            all_plan_group_data.append({
                "plan_idx": plan_idx,
                "plan_name": pname,
                "group_name": "未分組",
                "group_score": "-",
                "member_count": len(unassigned_members),
                "members": [f"{m['student_number']}號 {m['name']}" for m in unassigned_members]
            })

    # Header with separate member columns
    headers_g = ["分組模式", "小組名稱", "小組獨立得分", "組內成員人數"] + [f"成員 {i}" for i in range(1, max_members + 1)]
    ws_groups.append(headers_g)

    for item in all_plan_group_data:
        row = [
            item["plan_name"],
            item["group_name"],
            item["group_score"],
            item["member_count"]
        ] + item["members"]
        ws_groups.append(row)

        # Apply distinct pastel fill for this grouping plan
        row_idx = ws_groups.max_row
        row_fill = pastel_fills[item["plan_idx"] % len(pastel_fills)]
        for col_idx in range(1, len(row) + 1):
            cell = ws_groups.cell(row=row_idx, column=col_idx)
            cell.fill = row_fill
            cell.border = thin_border
            if col_idx in [3, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # --- Sheet 4: 質性紀錄明細 ---
    ws_notes = wb.create_sheet(title="質性紀錄明細")
    ws_notes.append([f"課程：{course_name} - 特殊表現與質性紀錄明細"])
    ws_notes.append([])
    ws_notes.append(["日期時間", "座號", "學生姓名", "紀錄內容"])

    notes_query = """
        SELECT qn.date, qn.timestamp, s.student_number, s.name AS student_name, qn.note_text
        FROM qualitative_notes qn
        JOIN students s ON qn.student_id = s.id
        WHERE qn.course_id = ?
    """
    notes_params = [course_id]
    if start_date:
        notes_query += " AND qn.date >= ?"
        notes_params.append(start_date)
    if end_date:
        notes_query += " AND qn.date <= ?"
        notes_params.append(end_date)
    notes_query += " ORDER BY qn.timestamp DESC"

    cursor.execute(notes_query, notes_params)
    for note in cursor.fetchall():
        ws_notes.append([
            note["timestamp"],
            note["student_number"],
            note["student_name"],
            note["note_text"]
        ])

    # Format styling for all sheets
    for ws in [ws_att, ws_score, ws_groups, ws_notes]:
        for row in ws.iter_rows():
            for cell in row:
                if cell.row == 3:  # Table header
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif cell.row == 1:
                    cell.font = title_font
                if cell.value is not None and cell.row > 3 and ws != ws_groups:
                    cell.border = thin_border

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"classroom_report_{course_id}_{get_today_str_taipei()}.xlsx"
    headers = {'Content-Disposition': f'attachment; filename="{filename}"'}
    return StreamingResponse(output, headers=headers, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
